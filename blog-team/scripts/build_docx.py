#!/usr/bin/env python3
"""Build article.docx for sign-off from an article folder.

Usage: python3 blog-team/scripts/build_docx.py articles/<slug>

Reads final.md (Markdown with YAML front matter) and, if present,
publish-package.json. Writes article.docx into the same folder: a sign-off
sheet (meta, categories, CTA, takeaways, open items) followed by the article.
Image slots, [[NEEDS: ...]] and [[LINK NEEDED: ...]] markers are highlighted.
"""
import json
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|\[\[[^\]]+\]\]|\[[^\]]+\]\([^)]+\))")


def split_front_matter(text):
    if text.startswith("---"):
        end = text.find("\n---", 3)
        if end != -1:
            try:
                import yaml
                meta = yaml.safe_load(text[3:end]) or {}
            except Exception:
                meta = {}
            return meta, text[end + 4:].lstrip("\n")
    return {}, text


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    r_id = part.relate_to(
        url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True
    )
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1F5C99")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(color)
    props.append(underline)
    run.append(props)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def add_inline(paragraph, text, site="https://weddedwonderland.com"):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("[[") and part.endswith("]]"):
            run = paragraph.add_run(part)
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
            run.bold = True
        elif part.startswith("[") and "](" in part:
            label, url = re.match(r"\[([^\]]+)\]\(([^)]+)\)", part).groups()
            add_hyperlink(paragraph, label, site + url if url.startswith("/") else url)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            paragraph.add_run(part[1:-1]).italic = True
        else:
            paragraph.add_run(part)


def add_body(doc, body):
    for raw in body.splitlines():
        line = raw.rstrip()
        if not line.strip():
            continue
        m = re.match(r"^(#{1,4})\s+(.*)", line)
        if m:
            doc.add_heading(m.group(2).strip(), level=len(m.group(1)) - 1 if len(m.group(1)) > 1 else 0)
            continue
        m = re.match(r"^\[IMAGE:\s*(.+?)\]$", line.strip())
        if m:
            p = doc.add_paragraph()
            run = p.add_run(f"[ IMAGE SLOT: {m.group(1)} ]")
            run.bold = True
            run.font.color.rgb = RGBColor(0x8A, 0x6D, 0x3B)
            run.font.highlight_color = WD_COLOR_INDEX.GRAY_25
            continue
        m = re.match(r"^\s*[-*]\s+(.*)", line)
        if m:
            add_inline(doc.add_paragraph(style="List Bullet"), m.group(1))
            continue
        m = re.match(r"^\s*\d+[.)]\s+(.*)", line)
        if m:
            add_inline(doc.add_paragraph(style="List Number"), m.group(1))
            continue
        if line.startswith(">"):
            p = doc.add_paragraph(style="Intense Quote")
            add_inline(p, line.lstrip("> "))
            continue
        add_inline(doc.add_paragraph(), line.strip())


def kv(doc, key, value):
    p = doc.add_paragraph()
    p.add_run(f"{key}: ").bold = True
    add_inline(p, str(value))


def main(folder):
    folder = Path(folder)
    meta, body = split_front_matter((folder / "final.md").read_text(encoding="utf-8"))
    pkg_path = folder / "publish-package.json"
    pkg = {}
    if pkg_path.exists():
        data = json.loads(pkg_path.read_text(encoding="utf-8"))
        pkg = data.get("payload", data)

    def get(key, default=""):
        return pkg.get(key) or meta.get(key) or default

    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    doc.add_heading("Sign-off sheet", level=1)
    kv(doc, "Title", get("title"))
    kv(doc, "Slug / URL", pkg.get("url") or get("slug"))
    kv(doc, "SEO title", get("seo_title"))
    kv(doc, "Meta description", get("meta_description"))
    kv(doc, "Excerpt", get("excerpt"))
    kv(doc, "Primary keyword", get("primary_keyword") or meta.get("primary_keyword", ""))
    cats = get("categories", [])
    kv(doc, "Primary category", pkg.get("primary_category") or (cats[0] if cats else ""))
    kv(doc, "All categories (first = primary)", ", ".join(cats))
    kv(doc, "Tags", ", ".join(get("tags", [])))
    kv(doc, "Author", get("author", "[[NEEDS: named author]]"))

    cta = pkg.get("cta") or {}
    if cta:
        doc.add_heading("CTA", level=2)
        for k in ("id", "heading", "body", "button", "url"):
            if cta.get(k):
                kv(doc, k.capitalize(), cta[k])

    takeaways = get("key_takeaways", [])
    if takeaways:
        doc.add_heading("Key takeaways", level=2)
        for t in takeaways:
            add_inline(doc.add_paragraph(style="List Bullet"), t)

    for label, key in (("Prices for Nic to sign off", "nic_signoff"), ("Open items (NEEDS)", "needs")):
        items = pkg.get(key) or []
        if items:
            doc.add_heading(label, level=2)
            for item in items:
                text = item if isinstance(item, str) else "; ".join(f"{k}: {v}" for k, v in item.items())
                add_inline(doc.add_paragraph(style="List Bullet"), text)

    images = pkg.get("images") or []
    if images:
        doc.add_heading("Images", level=2)
        table = doc.add_table(rows=1, cols=4)
        table.style = "Light Grid Accent 1"
        for cell, h in zip(table.rows[0].cells, ("Slot", "File", "Alt text", "Caption / credit")):
            cell.text = h
        for img in images:
            row = table.add_row().cells
            row[0].text = str(img.get("slot", ""))
            row[1].text = str(img.get("file_name", ""))
            row[2].text = str(img.get("alt_text", ""))
            row[3].text = f"{img.get('caption', '')} {img.get('credit', '')}".strip()

    doc.add_page_break()
    add_body(doc, body)

    faq = get("faq", [])
    if faq and not re.search(r"^#+\s*(FAQ|Frequently|Questions)", body, re.M | re.I):
        doc.add_heading("Frequently asked questions", level=2)
        for item in faq:
            doc.add_paragraph().add_run(item.get("q", "")).bold = True
            add_inline(doc.add_paragraph(), item.get("a", ""))

    out = folder / "article.docx"
    doc.save(out)
    print(f"Wrote {out}")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit("Usage: build_docx.py articles/<slug>")
    main(sys.argv[1])
